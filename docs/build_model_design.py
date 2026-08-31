from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("CRE_Expected_Loss_Model_Design.docx")
NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "5B6573"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
WHITE = "FFFFFF"
BLACK = "111111"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=11, bold=False, color=BLACK, italic=False, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (fld_char, instr, separate, text, end):
        run._r.append(item)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def add_callout(doc: Document, label: str, text: str, fill=PALE_BLUE) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), BLUE)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9.5, bold=True, color=NAVY)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value
            for run in cells[i].paragraphs[0].runs:
                set_font(run, size=9.5)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    set_table_geometry(table, widths)
    return table


def add_section_break(doc: Document) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
configure_styles(doc)

# Running header and footer.
header = section.header.paragraphs[0]
header.text = "CRE EXPECTED LOSS MODEL GOVERNANCE & DESIGN  |  SR 26-2 ALIGNED"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    set_font(run, size=8.5, bold=True, color=GRAY)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("Confidential working draft  |  Page ")
set_font(r, size=8.5, color=GRAY)
add_page_field(footer)

# Memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("MODEL GOVERNANCE & DESIGN")
set_font(r, size=23, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("Commercial Real Estate Expected Loss Framework | SR 26-2 Alignment")
set_font(r, size=15, color=GRAY)

metadata = [
    ("Document status", "Controlled working draft for model risk review"),
    ("Version", "1.0 draft"),
    ("As of date", "August 31, 2026"),
    ("Primary purpose", "Economic expected-loss measurement and scenario analysis"),
    ("Proposed grain", "Loan-property-quarter"),
    ("Implementation", "Python, Jupyter, DuckDB, and Parquet"),
    ("Controlling guidance", "SR 26-2, Revised Guidance on Model Risk Management (April 17, 2026)"),
    ("Legacy continuity", "SR 11-7 crosswalk retained; SR 11-7 is superseded"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True)
    r = p.add_run(value)
    set_font(r)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(14)
p_pr = p._p.get_or_add_pPr()
p_bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "10")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), BLUE)
p_bdr.append(bottom)
p_pr.append(p_bdr)

add_callout(
    doc,
    "Regulatory posture",
    "SR 26-2 superseded and replaced SR 11-7 on April 17, 2026. This document is structured as the model-specific governing record under SR 26-2 while preserving a legacy SR 11-7 crosswalk. It supports alignment but does not, by itself, establish compliance; operating evidence, approvals, effective challenge, monitoring, issue remediation, and audit remain necessary.",
)

add_heading(doc, "Document control and review", 1)
add_table(
    doc,
    ["Role", "Name", "Responsibility", "Status"],
    [
        ["Model owner", "TBD", "Purpose, use, assumptions, and ongoing performance", "Open"],
        ["Model developer", "TBD", "Data, methodology, code, testing, and documentation", "Open"],
        ["Independent validator", "TBD", "Conceptual soundness, outcomes analysis, limitations", "Open"],
        ["Approver", "TBD", "Scope, material assumptions, limitations, and permitted use", "Open"],
    ],
    [1700, 1300, 4560, 1800],
)

add_section_break(doc)
add_heading(doc, "1. Executive summary", 1)
doc.add_paragraph(
    "This document defines the proposed architecture for a Commercial Real Estate (CRE) expected-loss model built primarily from publicly available data. The model decomposes loss into probability of default (PD), loss given default (LGD), and exposure at default (EAD), applies forward-looking economic scenarios, and aggregates results across loans, properties, time horizons, and portfolio segments."
)
doc.add_paragraph(
    "This is the overarching model-specific document for design, use, controls, validation planning, monitoring, limitations, and governance. It should be read with the institution's enterprise model risk management policy, model inventory, validation reports, approvals, implementation evidence, monitoring reports, issue log, and internal-audit work. Those records are incorporated by reference through the evidence register in Appendix D."
)
add_callout(
    doc,
    "Use restriction",
    "Until the approval record is complete, outputs remain research estimates and may not be used as the sole basis for financial reporting, allowance, capital, underwriting, pricing, foreclosure, or regulatory submissions.",
    PALE_GOLD,
)
add_heading(doc, "1.1 Core equation", 2)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("EL(i,t) = sum over h [ marginal PD(i,t,h) x LGD(i,t,h) x EAD(i,t,h) x DF(h) ]")
set_font(r, size=12, bold=True, color=NAVY)
doc.add_paragraph(
    "Here, i identifies the exposure, t is the reporting date, h is the forecast interval, and DF is the selected discount factor. Marginal PD is conditional on survival to the start of each interval so that lifetime probabilities are not double-counted."
)
add_heading(doc, "1.2 Design principles", 2)
for item in [
    "Traceability: every reported result can be traced to source data, transformations, configuration, model version, and scenario.",
    "Separation of concerns: ingestion, feature engineering, estimation, scenario application, aggregation, validation, and reporting remain distinct.",
    "Determinism: authoritative loss calculations do not depend on an LLM or autonomous agent.",
    "Transparency before complexity: benchmarks and interpretable specifications precede nonlinear or highly parameterized models.",
    "Time-aware validation: development and testing respect reporting chronology and avoid future-information leakage.",
    "Conservative use of proxies: public-data gaps are disclosed, quantified where possible, and not disguised as observed loan facts.",
]:
    add_bullet(doc, item)

add_heading(doc, "2. Purpose, scope, and permitted use", 1)
add_heading(doc, "2.0 Regulatory basis and applicability", 2)
doc.add_paragraph(
    "SR 26-2 is the controlling interagency supervisory guidance for this document. It emphasizes a risk-based approach tailored to a banking organization's model risk profile and the size and complexity of its operations. It is expected to be most relevant to banking organizations with more than $30 billion in total assets, while potentially remaining relevant to smaller organizations with significant model risk exposure. Applicability to a particular institution must be confirmed by its legal, compliance, and model risk functions."
)
doc.add_paragraph(
    "Under SR 26-2, a model is a complex quantitative method, system, or approach underpinned by statistical, economic, or financial theory that processes input data into quantitative estimates. The PD, LGD, EAD, calibration, and scenario components meet this working classification. Simple arithmetic aggregation and deterministic rule-based software are controlled components but may fall outside the SR 26-2 model definition."
)
add_heading(doc, "2.1 Intended uses", 2)
for item in [
    "Estimate baseline and scenario-conditioned expected loss for CRE portfolios or representative segments.",
    "Compare risk across property types, geographies, vintages, structures, leverage, and refinance profiles.",
    "Evaluate sensitivities to net operating income (NOI), capitalization rates, vacancy, interest rates, and collateral values.",
    "Support research, model-development experiments, monitoring, benchmarking, and management discussion.",
]:
    add_bullet(doc, item)
add_heading(doc, "2.2 Out-of-scope uses for version 0.1", 2)
for item in [
    "Final financial-statement allowance or regulatory submission without a separately approved policy implementation.",
    "Automated credit approval, pricing, appraisal, foreclosure, or servicing action.",
    "Loan-level conclusions where required fields are imputed only from broad public proxies.",
    "Use outside the populations, property types, geographies, or economic regimes supported by development evidence.",
]:
    add_bullet(doc, item)
add_callout(doc, "Control boundary", "Results must be labeled research estimates until model validation, data controls, governance, and permitted-use approvals are complete.", PALE_GOLD)

add_heading(doc, "3. Modeling unit and portfolio segmentation", 1)
add_heading(doc, "3.1 Observation grain", 2)
doc.add_paragraph(
    "The preferred observation is one loan-property-quarter. A loan secured by multiple properties should be represented through documented allocation rules or retained at loan-quarter grain with collateral-pool features. A property supporting multiple loans should retain lien position and allocated value. Duplicate economic exposure must not be introduced through joins."
)
add_heading(doc, "3.2 Minimum keys", 2)
add_table(
    doc,
    ["Key", "Purpose", "Required control"],
    [
        ["loan_id", "Stable exposure identifier", "Unique within source and mapped across reporting periods"],
        ["property_id", "Stable collateral identifier", "Supports multi-loan and multi-property relationships"],
        ["reporting_date", "Observation timestamp", "Quarter-end normalization and no future joins"],
        ["source_system", "Lineage", "Preserved from landing through reporting"],
        ["record_version", "Restatement handling", "Latest-valid rule with prior versions retained"],
    ],
    [1700, 3000, 4660],
)
add_heading(doc, "3.3 Segmentation", 2)
doc.add_paragraph("Core segmentation should be explicit and stable enough for monitoring:")
for item in [
    "Property type: office, retail, multifamily, industrial, hotel, healthcare, self-storage, mixed-use, and other.",
    "Geography: metropolitan area where reliable; otherwise state, Census division, or liquidity tier.",
    "Loan structure: fixed/floating, amortizing/interest-only, balloon, recourse, lien position, construction/stabilized.",
    "Risk state: current, watchlist, delinquent, nonaccrual, modified, defaulted, workout, resolved.",
    "Vintage and horizon: origination cohort, remaining term, maturity concentration, and refinance window.",
]:
    add_bullet(doc, item)

add_heading(doc, "4. Data architecture", 1)
add_heading(doc, "4.1 Data layers", 2)
add_number(doc, "Raw: immutable source extracts plus retrieval date, URL, license, checksum, and schema snapshot.")
add_number(doc, "Interim: normalized dates, identifiers, units, categories, and source-specific quality flags.")
add_number(doc, "Processed: point-in-time loan-property-quarter model table and separately versioned targets.")
add_number(doc, "Outputs: predictions, scenarios, diagnostics, aggregations, and run manifests.")
add_heading(doc, "4.2 Candidate public-data families", 2)
add_table(
    doc,
    ["Data family", "Illustrative use", "Primary limitation"],
    [
        ["Public CMBS disclosures", "Loan, property, performance, maturity, delinquency, and workout evidence", "Coverage and field consistency vary by deal and reporting period"],
        ["Regulatory bank data", "Portfolio benchmarks, charge-offs, delinquency, and concentration context", "Usually aggregated rather than loan level"],
        ["Macroeconomic series", "Rates, unemployment, inflation, output, and scenario drivers", "National series may be weak proxies for local CRE conditions"],
        ["Property-market series", "Rents, vacancy, cap rates, prices, permits, and transaction liquidity", "Licensing, coverage, and methodology can differ materially"],
        ["Geographic reference data", "Metro mapping, local employment, demographics, and disaster exposure", "Temporal joins and boundary changes require care"],
    ],
    [2100, 3400, 3860],
)
add_heading(doc, "4.3 Point-in-time controls", 2)
for item in [
    "Every feature must carry an availability date distinct from its economic reference date.",
    "Revised macro series must be versioned or replaced by vintage data when the use case requires real-time replication.",
    "Outcomes and workout cash flows must be isolated from the feature snapshot used to predict them.",
    "Train, validation, and test partitions must be based on time, with entity grouping where leakage is possible.",
]:
    add_bullet(doc, item)

add_heading(doc, "5. Outcome definitions", 1)
add_heading(doc, "5.1 Default", 2)
doc.add_paragraph(
    "The proposed default event is the earliest observable occurrence of a material credit failure. Candidate triggers are 90 or more days past due, nonaccrual, foreclosure, transfer to real-estate-owned status, a distressed restructuring with economic concession, or a realized credit loss. Source-specific indicators should map into one canonical event table while retaining the original trigger."
)
add_heading(doc, "5.2 Cures and repeated defaults", 2)
doc.add_paragraph(
    "A cure does not erase the original default event. For first-default PD development, only the first qualifying default is used. Re-default analysis should be developed separately with an explicit cure window and risk-entry date. Technical or reporting-only delinquency reversals should be handled through documented exception rules."
)
add_heading(doc, "5.3 LGD outcome", 2)
doc.add_paragraph(
    "Realized LGD should use discounted net workout cash flows from default through resolution. Recoveries include principal collections, collateral sale proceeds, guarantees, and other credit support. Costs include legal, servicing, preservation, taxes, disposition, and other directly attributable workout expenses."
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("LGD = 1 - [ present value of net recoveries / EAD at default ]")
set_font(r, size=12, bold=True, color=NAVY)
doc.add_paragraph(
    "Unresolved defaults must not be treated as zero loss. They require censoring-aware methods, a resolution model, or a documented conservative estimate. Recoveries above EAD may be capped or retained for analysis according to approved policy."
)

add_heading(doc, "6. Feature framework", 1)
add_table(
    doc,
    ["Feature family", "Representative variables", "Main risks"],
    [
        ["Leverage", "Current/origination LTV, stressed LTV, debt yield", "Appraisal staleness and valuation circularity"],
        ["Cash flow", "NOI, DSCR, NOI growth, occupancy, rent rollover", "Inconsistent definitions and annualization"],
        ["Structure", "Rate type, IO period, amortization, lien, recourse", "Missing contractual detail"],
        ["Refinance", "Remaining term, balloon share, rate reset, refinance DSCR", "Future-rate and proceeds assumptions"],
        ["Market", "Vacancy, rent growth, cap rate, price index, liquidity", "Proxy geography and licensing constraints"],
        ["Behavior", "Delinquency transitions, modifications, watchlist, servicer transfer", "Potential intervention and reporting bias"],
        ["Macro", "Rates, unemployment, GDP, inflation, credit spreads", "Frequency mismatch and multicollinearity"],
    ],
    [1800, 3980, 3580],
)
add_heading(doc, "6.1 Feature policies", 2)
for item in [
    "Transformations are fitted on training data only and versioned with the model.",
    "Missingness indicators are retained when absence may be informative; imputation must be documented by feature and segment.",
    "Winsorization or clipping thresholds are estimated without test-period information.",
    "Property values used in stressed LTV must distinguish observed appraisal, indexed value, and scenario-implied value.",
    "Features with direct post-default information are prohibited from pre-default PD development.",
]:
    add_bullet(doc, item)

add_heading(doc, "7. PD methodology", 1)
add_heading(doc, "7.1 Recommended structure", 2)
doc.add_paragraph(
    "Use a quarterly discrete-time hazard model as the primary PD architecture. Each surviving exposure contributes one row per quarter until default, payoff, maturity, censoring, or the maximum horizon. The model estimates conditional default hazard; marginal PD is hazard multiplied by survival through prior intervals."
)
add_heading(doc, "7.2 Development ladder", 2)
add_number(doc, "Portfolio and segment historical-rate benchmarks with credible intervals.")
add_number(doc, "Regularized logistic discrete-time hazard model with interpretable transformations.")
add_number(doc, "Nonlinear challenger, such as gradient boosting, only after leakage, stability, and calibration controls are established.")
add_number(doc, "Calibration by horizon and material segment using out-of-time evidence.")
add_heading(doc, "7.3 Evaluation", 2)
for item in [
    "Discrimination: ROC AUC, precision-recall AUC, rank ordering, and concentration of defaults.",
    "Calibration: observed-to-expected ratios, reliability curves, Brier score, and calibration slope/intercept.",
    "Stability: feature drift, score drift, segment performance, vintage performance, and economic-regime sensitivity.",
    "Survival consistency: nonnegative hazards, probabilities bounded by one, and cumulative PD monotonic with horizon.",
]:
    add_bullet(doc, item)

add_heading(doc, "8. LGD methodology", 1)
doc.add_paragraph(
    "LGD should combine a workout-based empirical model with a collateral-value benchmark. A two-stage specification is preferred when zero or near-zero loss outcomes are material: first estimate the probability of positive loss, then estimate severity conditional on positive loss. Alternative bounded-loss or fractional-response models may be compared."
)
add_heading(doc, "8.1 Principal drivers", 2)
for item in [
    "LTV and stressed LTV at default; lien position and recourse.",
    "Property type, market liquidity, cap rate, occupancy, NOI trajectory, and transaction conditions.",
    "Time to resolution, workout path, legal environment, and expenses.",
    "Guarantee or credit support, servicer strategy, and collateral complexity.",
]:
    add_bullet(doc, item)
add_heading(doc, "8.2 Downturn treatment", 2)
doc.add_paragraph(
    "Scenario-conditioned LGD should react to both collateral value and workout timing. Avoid applying an arbitrary additive overlay when the same macro shock is already transmitted through NOI, cap rates, liquidity, or resolution duration. Any management overlay must be separately identified, justified, sensitivity-tested, approved, and reversible."
)

add_heading(doc, "9. EAD methodology", 1)
doc.add_paragraph(
    "For fully funded term loans, EAD begins with the contractual balance projected to each horizon using scheduled amortization, prepayment assumptions, modifications, and maturity treatment. For revolving, construction, or future-funding commitments, EAD includes expected additional draws through a calibrated credit conversion factor or utilization model."
)
add_heading(doc, "9.1 EAD controls", 2)
for item in [
    "Reconcile projected balances to contractual schedules and source-system balances.",
    "Prevent negative balances and funding beyond commitment limits.",
    "Separate scheduled maturity from default; define refinance failure behavior explicitly.",
    "Avoid double counting accrued interest, fees, guarantees, or unfunded exposure across EAD and LGD.",
]:
    add_bullet(doc, item)

add_heading(doc, "10. Scenario framework", 1)
doc.add_paragraph(
    "Scenarios should translate coherent macro and property-market paths into loan-level cash flow, value, refinancing, PD, LGD, and EAD effects. Baseline, adverse, and severely adverse labels are descriptive; the actual variable paths and weights require approval for each use case."
)
add_heading(doc, "10.1 Transmission chain", 2)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Macro path -> rent / vacancy / NOI -> value and refinance capacity -> PD, LGD, EAD -> expected loss")
set_font(r, size=11.5, bold=True, color=NAVY)
add_heading(doc, "10.2 Scenario mechanics", 2)
for item in [
    "Property type and geography determine sensitivity, not just portfolio-average shocks.",
    "Cap-rate shocks and NOI shocks jointly determine scenario-implied collateral value.",
    "Floating-rate debt service and refinance coupons follow scenario rate paths with contractual lags and floors.",
    "Scenario weights sum to one and are stored with effective date, rationale, owner, and approval.",
    "Weighted results retain each scenario separately so aggregation can be reproduced.",
]:
    add_bullet(doc, item)

add_heading(doc, "11. Expected-loss aggregation", 1)
doc.add_paragraph(
    "The calculation engine consumes horizon-level marginal PD, LGD, EAD, and discount factors. It produces loan-horizon expected loss before aggregation. Scenario-weighted loss is calculated only after scenario-specific results are preserved."
)
add_table(
    doc,
    ["Control", "Requirement"],
    [
        ["Probability", "PD and LGD lie within approved bounds; cumulative PD does not decline with horizon"],
        ["Exposure", "EAD is finite, nonnegative, and reconciled to balances plus permissible commitments"],
        ["Discounting", "Rate basis, timing convention, and compounding match the selected framework"],
        ["Scenario weights", "Weights are nonnegative, sum to one, and match the run manifest"],
        ["Reconciliation", "Loan-horizon results sum exactly to scenario, segment, and portfolio totals"],
    ],
    [2300, 7060],
)

add_heading(doc, "12. Validation framework", 1)
doc.add_paragraph(
    "Validation rigor, timing, nature, and frequency are risk-based and must align with model approach, use, materiality, methodology, change frequency, and data limitations. Validation ordinarily occurs before first use. Any temporary pre-validation use requires documented urgency, communicated limitations, compensating controls, restricted use, enhanced monitoring, an accountable approver, and a dated completion plan."
)
add_heading(doc, "12.1 Conceptual soundness", 2)
for item in [
    "Assess outcome definitions, segmentation, feature rationale, assumptions, estimator choice, calibration, and scenario transmission.",
    "Compare primary models with simple benchmarks and credible alternative specifications.",
    "Evaluate data representativeness and the effect of public-data limitations on intended use.",
]:
    add_bullet(doc, item)
add_heading(doc, "12.2 Implementation and process verification", 2)
doc.add_paragraph(
    "SR 26-2 does not retain process verification as a separately named validation component, but implementation verification remains an essential internal control and a source of evidence supporting reliability. It is retained here as an explicit requirement."
)
for item in [
    "Independent code review and unit tests for formulas, joins, time alignment, scenario weighting, and aggregation.",
    "Reproducible runs from immutable inputs, configuration, dependency lock, and source-control commit.",
    "Boundary, missing-value, duplicate-key, and adversarial data tests.",
]:
    add_bullet(doc, item)
add_heading(doc, "12.3 Outcomes analysis", 2)
for item in [
    "Out-of-time and, where possible, out-of-segment performance.",
    "Backtesting by property type, geography, vintage, risk state, and forecast horizon.",
    "Sensitivity and stress testing, including correlated shocks and missing-data alternatives.",
    "Benchmark comparison, override analysis, and stability monitoring with thresholds and escalation.",
]:
    add_bullet(doc, item)

add_heading(doc, "13. Model risk, limitations, and safeguards", 1)
add_table(
    doc,
    ["Risk", "Potential effect", "Safeguard"],
    [
        ["Public-data selection bias", "CMBS or disclosed loans may not represent bank-held CRE", "Population comparison, segmentation, reweighting tests, and explicit scope limits"],
        ["Sparse defaults", "Unstable parameters and misleading segment estimates", "Pooling, regularization, shrinkage, benchmarks, and uncertainty intervals"],
        ["Stale collateral values", "Understated leverage and LGD", "Value timestamp controls, indexed and scenario-implied alternatives"],
        ["Economic regime change", "Historical relationships may not persist", "Rolling monitoring, stress tests, recalibration triggers, and expert challenge"],
        ["Outcome censoring", "Unresolved workouts bias LGD downward", "Censoring-aware treatment and maturity/resolution diagnostics"],
        ["Automation overreach", "Unreviewed outputs influence decisions", "Permitted-use controls, human approval, logs, and deterministic calculations"],
    ],
    [1900, 3100, 4360],
)

add_heading(doc, "14. Technology and implementation architecture", 1)
add_heading(doc, "14.1 Core stack", 2)
for item in [
    "Python modules for transformations, estimation, scoring, validation, and reporting.",
    "Jupyter notebooks for exploration, development evidence, and reviewable analysis.",
    "DuckDB for local analytical queries and controlled dataset materialization.",
    "Parquet for typed, compressed, columnar interchange between layers.",
    "YAML configuration for paths, definitions, horizons, thresholds, scenarios, and approved parameters.",
    "Git for source code and documentation; large datasets remain outside ordinary source control.",
]:
    add_bullet(doc, item)
add_heading(doc, "14.2 Optional application layers", 2)
doc.add_paragraph(
    "A web interface may later be implemented with Flask or another framework. LangGraph or CrewAI may orchestrate review, research, or narrative workflows, but they should not replace deterministic calculations or governance controls. SR 26-2 states that generative and agentic AI models are outside its scope; nevertheless, institutional risk-management and governance practices should determine controls for such tools. Agent-generated commentary must be labeled, sourced, reviewable, and separable from model output."
)
add_heading(doc, "14.3 Proposed run manifest", 2)
add_table(
    doc,
    ["Field", "Example content"],
    [
        ["run_id", "Unique immutable identifier"],
        ["as_of_date", "Portfolio reporting date"],
        ["input_versions", "Checksums or dataset snapshot identifiers"],
        ["code_version", "Git commit hash"],
        ["model_versions", "PD, LGD, EAD, calibration, and scenario versions"],
        ["configuration", "Resolved configuration hash"],
        ["execution", "Start/end time, environment, warnings, and status"],
        ["approvals", "Reviewer, date, decision, and limitations"],
    ],
    [2500, 6860],
)

add_heading(doc, "15. Governance and change management", 1)
for item in [
    "Material changes include outcome definitions, population, horizon, methodology, scenario transmission, calibration, overlays, or permitted use.",
    "Each material change requires documented rationale, impact analysis, testing, independent challenge, approval, and version increment.",
    "Emergency overrides must be time-limited, separately reported, approved, and subject to retrospective review.",
    "Monitoring thresholds must specify warning, breach, escalation owner, remediation deadline, and use restriction where appropriate.",
    "Model artifacts must include development evidence, validation findings, implementation tests, limitations, user guidance, and change history.",
]:
    add_bullet(doc, item)

add_heading(doc, "16. Open design decisions", 1)
doc.add_paragraph("The following decisions should be resolved before version 1.0 implementation is considered complete.")
add_table(
    doc,
    ["Decision", "Recommended starting position", "Approval needed"],
    [
        ["Authoritative framework", "Economic expected loss first; map CECL/IFRS 9 later", "Model owner and accounting/risk policy"],
        ["Forecast horizon", "Quarterly through contractual maturity with configurable cap", "Model owner"],
        ["Default definition", "Earliest material failure using canonical trigger hierarchy", "Credit policy and validation"],
        ["Discount rate", "Effective-rate or economic-rate convention selected by use case", "Accounting/risk policy"],
        ["Prepayment and maturity", "Separate competing exits and explicit refinance treatment", "Model owner and treasury/credit"],
        ["Scenario set and weights", "Three coherent scenarios with documented effective dates", "Governance committee"],
        ["Unresolved LGD cases", "Censoring-aware treatment plus conservative benchmark", "Model owner and validation"],
        ["Public-data suitability", "Permit segment analytics; restrict unsupported loan-level claims", "Model risk and business owner"],
        ["Production threshold", "Validation, controls, monitoring, and approvals completed", "Governance committee"],
    ],
    [2500, 4060, 2800],
)

add_heading(doc, "17. Phased delivery plan", 1)
add_number(doc, "Foundation: approve purpose, data contract, default/LGD definitions, lineage, and benchmark calculations.")
add_number(doc, "PD development: build hazard dataset, baseline model, challenger, calibration, and out-of-time evaluation.")
add_number(doc, "LGD/EAD development: construct workout and balance histories, benchmarks, models, and uncertainty treatment.")
add_number(doc, "Scenario integration: specify paths, transmission equations, weights, sensitivities, and reconciliation controls.")
add_number(doc, "Validation and governance: independent review, remediation, documentation, monitoring, and permitted-use approval.")
add_number(doc, "Application layer: only after stable calculations, add UI, automation, access controls, and deployment monitoring.")

add_heading(doc, "18. Model classification and risk assessment", 1)
doc.add_paragraph(
    "SR 26-2 frames overall model risk through inherent risk and materiality. Materiality reflects model exposure and purpose. The assessment is performed at least at initial classification, before first use, after material change, when use expands, and when monitoring indicates a change in risk."
)
add_table(
    doc,
    ["Dimension", "Assessment for this model", "Evidence / action"],
    [
        ["Model status", "Model: PD, LGD, EAD, calibration, and statistical scenario relationships", "Inventory entry and component map required"],
        ["Inherent risk", "Provisionally high due to multiple components, assumptions, public-data constraints, and scenario dependence", "Complete scored inherent-risk assessment"],
        ["Exposure", "TBD; quantify portfolio balance, allowance/capital influence, users, frequency, and downstream consumers", "Exposure inventory and dependency map"],
        ["Purpose", "Potential financial-risk management and possibly financial/regulatory reporting", "Approve intended and prohibited uses"],
        ["Materiality", "Provisionally high if used for allowance, capital, stress testing, or material portfolio decisions", "MRM approval of tier and rationale"],
        ["Aggregate risk", "Shared data, assumptions, scenarios, and downstream reports can create correlated error", "Enterprise dependency and concentration assessment"],
        ["Overall rating", "Provisional High pending institutional methodology", "Final classification by Model Risk Management"],
    ],
    [1900, 4320, 3140],
)
add_callout(
    doc,
    "Provisional classification",
    "Apply the institution's approved rating methodology. Until exposure and purpose are approved, manage the model at the higher provisional tier rather than using this document's narrative as the final rating.",
    PALE_GOLD,
)

add_heading(doc, "19. Lifecycle governance and accountability", 1)
add_heading(doc, "19.1 Lifecycle gates", 2)
add_number(doc, "Concept approval: intended use, model definition, provisional risk rating, data feasibility, and accountable owner.")
add_number(doc, "Development completion: theory, data, assumptions, testing, limitations, code review, documentation, and owner attestation.")
add_number(doc, "Validation completion: conceptual soundness, outcomes analysis, monitoring design, implementation verification, findings, and use recommendation.")
add_number(doc, "Use approval: model-risk decision, business acceptance, conditions, limitations, monitoring thresholds, and effective date.")
add_number(doc, "Ongoing use: scheduled monitoring, outcome comparison, limitation review, issue remediation, and periodic relevance assessment.")
add_number(doc, "Change or retirement: impact assessment, revalidation decision, controlled implementation, archival, and downstream transition.")
add_heading(doc, "19.2 Responsibilities", 2)
add_table(
    doc,
    ["Role", "Accountability"],
    [
        ["Board / delegated committee", "Oversees aggregate model risk within risk appetite and receives material risk reporting as defined by enterprise policy"],
        ["Senior management", "Provides resources, enforces policy, reviews material limitations/findings, and directs remediation"],
        ["Model owner", "Accountable for intended use, performance, inventory accuracy, monitoring, limitations, issues, and user communication"],
        ["Developer", "Documents design, data, assumptions, testing, code, reproducibility, and changes; responds to effective challenge"],
        ["Model user", "Uses only for approved purposes, understands limitations, reports unexpected behavior, and preserves decision evidence"],
        ["Model Risk Management / validator", "Provides objective, technically competent challenge with sufficient independence and influence; recommends approval conditions"],
        ["Technology / data owner", "Maintains controlled implementation, access, lineage, change deployment, resilience, and production evidence"],
        ["Internal audit", "Evaluates whether the model-risk framework and related policies are rigorous, effective, and implemented; does not duplicate development or validation"],
    ],
    [2500, 6860],
)

add_heading(doc, "20. Approval, conditions, and use control", 1)
add_table(
    doc,
    ["Approval field", "Required record"],
    [
        ["Model identifier / version", "TBD - match the authoritative model inventory"],
        ["Risk rating / tier", "TBD - approved under institutional methodology"],
        ["Approved uses", "TBD - enumerate by process, portfolio, legal entity, geography, and reporting purpose"],
        ["Prohibited uses", "Sole-decision automation; unsupported populations; unapproved accounting, capital, underwriting, or regulatory use"],
        ["Approval conditions", "TBD - validation findings, overlays, limits, enhanced monitoring, or data remediation"],
        ["Effective / expiry dates", "TBD - include next review or revalidation trigger"],
        ["Approvers", "Business owner, Model Risk Management, and other functions required by enterprise policy"],
        ["Evidence location", "Authoritative approval record and electronic signatures"],
    ],
    [2600, 6760],
)
add_heading(doc, "20.1 Temporary use before validation", 2)
doc.add_paragraph(
    "Temporary use before completed validation is exceptional. The request must document urgent need, scope, duration, known limitations, affected decisions, exposure, compensating controls, heightened monitoring, user notices, approval authority, validation completion date, and automatic expiration. It must not become an informal permanent exception."
)

add_heading(doc, "21. Ongoing monitoring and performance response", 1)
doc.add_paragraph(
    "The owner maintains a monitoring plan approved with the model. Thresholds are set from developmental and validation evidence and are recorded numerically in the monitoring specification; this document provides the minimum categories rather than inventing institution-specific limits."
)
add_table(
    doc,
    ["Monitoring domain", "Minimum measures", "Response to breach"],
    [
        ["Data", "Completeness, timeliness, key reconciliation, missingness, range, drift, and source changes", "Investigate lineage; assess affected runs; restrict use if material"],
        ["PD", "Calibration, O/E, rank ordering, Brier score, horizon consistency, segment and vintage stability", "Recalibrate, adjust, redevelop, or constrain use"],
        ["LGD", "Recovery timing, realized severity, resolution mix, collateral error, segment bias", "Re-estimate, overlay with governance, or redevelop"],
        ["EAD", "Balance reconciliation, utilization error, maturity/prepayment performance, exposure bounds", "Correct implementation or assumptions; restate affected output"],
        ["Scenario", "Sensitivity, directionality, coherence, weight controls, duplicate-shock tests", "Suspend scenario result or revise approved transmission"],
        ["Use and exposure", "Users, decisions, portfolio size, downstream dependencies, overrides, exceptions", "Reclassify risk/materiality and reassess controls"],
        ["Limitations/issues", "Aging, severity, overdue actions, repeated breaches, compensating controls", "Escalate; impose limits or withdraw approval"],
    ],
    [1900, 4520, 2940],
)
add_heading(doc, "21.1 Performance actions", 2)
doc.add_paragraph(
    "Meaningful deviation from expectations triggers a documented assessment of adjustment, recalibration, redevelopment, overlay, use restriction, or retirement. The action must identify root cause, materiality, affected outputs and decisions, interim control, owner, due date, approval, and closure evidence."
)

add_heading(doc, "22. Overlays, expert judgment, and exceptions", 1)
for item in [
    "Every overlay has a defined purpose, quantified basis, scope, owner, effective date, expiration or review date, and independent challenge.",
    "The rationale distinguishes model limitation, emerging risk, data deficiency, policy choice, and conservatism.",
    "Double counting is tested across base model, scenario effects, qualitative adjustments, and management overlays.",
    "Backtesting compares adjusted and unadjusted results; overlays do not substitute indefinitely for remediation.",
    "Exceptions to data, validation, use, or monitoring standards are time-bound, approved, tracked, and reported by severity.",
]:
    add_bullet(doc, item)

add_heading(doc, "23. Third-party and external-resource controls", 1)
doc.add_paragraph(
    "Public data, vendor data, external parameters, libraries, and any third-party model component remain subject to model-risk management. Limited access to proprietary code or development data does not remove the need to understand conceptual design, intended use, development data, performance, limitations, and local fit."
)
for item in [
    "Perform due diligence covering methodology, data provenance, versioning, change notices, performance evidence, limitations, service continuity, information security, and contractual audit/access rights.",
    "Validate local use, customization, data mapping, and performance; vendor validation does not replace institution-specific evaluation.",
    "Monitor outcomes and continued fitness for purpose; assess vendor releases as model changes.",
    "Integrate external validators or consultants into institutional issue, approval, documentation, and accountability processes.",
]:
    add_bullet(doc, item)

add_heading(doc, "24. Issue management and model-risk acceptance", 1)
add_table(
    doc,
    ["Issue field", "Minimum content"],
    [
        ["Identification", "Unique ID, source, date, model/version, component, and affected use"],
        ["Assessment", "Severity, root cause, exposure, affected outputs/decisions, and aggregate implications"],
        ["Interim control", "Use limit, overlay, manual review, enhanced monitoring, or suspension"],
        ["Remediation", "Action, accountable owner, milestones, due date, dependencies, and resources"],
        ["Governance", "Approver, risk acceptance if any, expiration, escalation, and reporting"],
        ["Closure", "Testing, independent verification where required, approval, and retained evidence"],
    ],
    [2300, 7060],
)

add_heading(doc, "25. Documentation hierarchy and evidence standard", 1)
doc.add_paragraph(
    "This document is the umbrella model record; it should not absorb every analysis or operating record. Supporting evidence remains separately versioned, reproducible, accessible to reviewers, and linked through stable identifiers. A requirement is not satisfied merely because this document states that an activity should occur."
)
add_table(
    doc,
    ["Artifact", "Purpose", "Status"],
    [
        ["Enterprise MRM policy and procedures", "Institution-wide definitions, risk rating, roles, approval, validation, monitoring, exceptions, and audit", "External dependency"],
        ["Authoritative model inventory record", "Identifier, owner, status, purpose, exposure, tier, dependencies, validation, issues, and dates", "To be completed"],
        ["Development report and reproducible notebooks", "Theory, data, assumptions, estimation, testing, alternatives, limitations, and results", "In development"],
        ["Implementation verification", "Code/data reconciliation, configuration, access, deployment, lineage, and run controls", "To be completed"],
        ["Validation report", "Objective challenge, findings, limitations, performance assessment, and use recommendation", "To be completed independently"],
        ["Approval record", "Approved/prohibited use, conditions, signatures, effective/expiry dates", "To be completed"],
        ["Monitoring specification and reports", "Metrics, thresholds, frequency, owners, breaches, actions, and trend", "To be completed"],
        ["Issue and exception log", "Findings, risk acceptance, remediation, aging, escalation, and closure", "To be completed"],
        ["Change log and release evidence", "Materiality assessment, testing, validation decision, approval, deployment, and rollback", "To be completed"],
        ["Internal-audit evidence", "Framework effectiveness and policy-implementation assessment", "Enterprise responsibility"],
    ],
    [2500, 4940, 1920],
)

add_heading(doc, "26. Regulatory alignment conclusion", 1)
doc.add_paragraph(
    "The design addresses SR 26-2's principal model-specific themes: risk-based tailoring; model definition and risk assessment; clear purpose and use; disciplined development and testing; limitations; conceptual soundness; outcomes analysis; ongoing monitoring; governance and controls; accountability and conflicts; inventory; documentation; and third-party products. Legacy SR 11-7 concepts remain traceable where they continue to strengthen practice, including implementation verification, benchmarks, effective challenge, inventory detail, and auditability."
)
add_callout(
    doc,
    "Readiness conclusion",
    "Documentation alignment is not operational compliance. The model is not ready for governed production use until the evidence register is populated, validation is completed with sufficient independence, findings and use conditions are resolved or accepted, approval is recorded, and monitoring operates as designed.",
    PALE_GOLD,
)

add_heading(doc, "Appendix A. Minimum model-ready schema", 1)
add_table(
    doc,
    ["Category", "Minimum fields"],
    [
        ["Identifiers", "loan_id, property_id, source_system, reporting_date, record_version"],
        ["Exposure", "balance, commitment, undrawn amount, interest rate, payment, maturity, amortization"],
        ["Collateral", "property type, location, value, value date, lien, occupancy, NOI, cap rate"],
        ["Performance", "delinquency, nonaccrual, modification, watchlist, servicer/workout status"],
        ["Outcomes", "default date/trigger, EAD at default, recoveries, expenses, cash-flow dates, resolution"],
        ["Market and macro", "series identifier, geography, reference date, availability date, vintage"],
        ["Lineage and quality", "source file, ingestion timestamp, transformation version, quality flags"],
    ],
    [2100, 7260],
)

add_heading(doc, "Appendix B. Acceptance criteria for the design phase", 1)
for item in [
    "Purpose, intended use, exclusions, population, horizon, and grain are approved.",
    "Default, cure, loss, recovery, expense, resolution, and EAD definitions are executable from available data.",
    "Point-in-time feature and outcome construction is demonstrated without leakage.",
    "Baseline PD, LGD, EAD, and expected-loss calculations reconcile on a controlled sample.",
    "Scenario transmission is mathematically specified and avoids duplicate shocks.",
    "Validation plan, monitoring thresholds, limitations, ownership, and change controls are documented.",
    "Repository structure, test strategy, configuration conventions, and run manifest are agreed.",
]:
    add_bullet(doc, item)

add_heading(doc, "Appendix C. SR 26-2 and legacy SR 11-7 crosswalk", 1)
add_table(
    doc,
    ["Guidance theme", "Document coverage", "Required operating evidence"],
    [
        ["SR 26-2 I-II: risk-based scope and tailoring", "Sections 1-2 and 18 define scope, applicability, provisional risk, purpose, exposure, and materiality", "Institutional applicability analysis and approved model-risk rating"],
        ["SR 26-2 III: individual and aggregate model risk; effective challenge", "Sections 13, 18, 19, and 24 address inherent risk, dependencies, challenge, issues, and influence", "Inventory dependencies, aggregate-risk reporting, challenge records, and issue escalation"],
        ["SR 26-2 IV: development and use", "Sections 3-11 define purpose, data, outcomes, features, methods, testing, use limits, and scenarios", "Development evidence, testing results, user procedures, and use attestations"],
        ["SR 26-2 V: conceptual soundness", "Sections 6-12 provide theory, assumptions, alternatives, benchmarks, and validation scope", "Completed validation workpapers and report"],
        ["SR 26-2 V: outcomes analysis", "Sections 7-12 and 21 specify backtesting, calibration, realized outcomes, and thresholds", "Outcomes datasets, reports, breach decisions, and remediation"],
        ["SR 26-2 V: ongoing monitoring", "Section 21 defines domains, responses, limitation review, and performance actions", "Approved numeric thresholds and recurring monitoring reports"],
        ["SR 26-2 VI: governance, roles, inventory, documentation", "Sections 15 and 18-25 define lifecycle gates, responsibilities, approvals, inventory, evidence, issues, and changes", "Policy, inventory record, approvals, committee minutes, issue log, and audit evidence"],
        ["SR 26-2 VII: vendor/third-party products", "Section 23 addresses due diligence, local validation, change monitoring, and external-resource oversight", "Due diligence, contracts, vendor evidence, local testing, and release reviews"],
        ["Legacy SR 11-7: development, implementation, and use", "Sections 3-11, 12.2, 14, and 17 retain point-in-time construction, implementation verification, and use controls", "Implementation test pack, reconciliations, production controls, and user records"],
        ["Legacy SR 11-7: comprehensive validation", "Section 12 retains conceptual soundness, outcomes analysis, monitoring, benchmarks, and process checks", "Validation report, workpapers, findings, and validation schedule"],
        ["Legacy SR 11-7: governance, policy, inventory, audit", "Sections 15 and 18-25 address roles, authority, documentation, inventory, issues, and audit linkage", "Enterprise policy, inventory, board/senior reporting, and internal-audit assessment"],
    ],
    [2150, 4050, 3160],
)

add_heading(doc, "Appendix D. Model evidence register", 1)
add_table(
    doc,
    ["Evidence ID", "Record", "Owner", "Location / version", "Status"],
    [
        ["E-01", "Enterprise MRM policy applicability", "MRM / Compliance", "TBD", "Open"],
        ["E-02", "Model inventory and risk rating", "Model Owner / MRM", "TBD", "Open"],
        ["E-03", "Data inventory, lineage, and quality report", "Data Owner", "TBD", "Open"],
        ["E-04", "Development report and test results", "Developer", "TBD", "Open"],
        ["E-05", "Implementation verification and reconciliation", "Technology / MRM", "TBD", "Open"],
        ["E-06", "Validation report and findings", "Validator", "TBD", "Open"],
        ["E-07", "Approval and use conditions", "MRM / Approver", "TBD", "Open"],
        ["E-08", "Monitoring specification and threshold approval", "Model Owner / MRM", "TBD", "Open"],
        ["E-09", "Recurring monitoring and outcomes reports", "Model Owner", "TBD", "Open"],
        ["E-10", "Issue, exception, and overlay log", "Model Owner / MRM", "TBD", "Open"],
        ["E-11", "Change and release log", "Developer / Technology", "TBD", "Open"],
        ["E-12", "Internal-audit assessment", "Internal Audit", "TBD", "Open"],
    ],
    [1100, 3000, 1900, 2260, 1100],
)

add_heading(doc, "Appendix E. Authoritative sources", 1)
p = doc.add_paragraph()
p.add_run("Controlling guidance: ").bold = True
add_hyperlink(p, "SR 26-2, Revised Guidance on Model Risk Management (April 17, 2026)", "https://www.federalreserve.gov/supervisionreg/srletters/SR2602.htm")
p = doc.add_paragraph()
p.add_run("Attachment: ").bold = True
add_hyperlink(p, "Supervisory Guidance on Model Risk Management (SR 26-2 attachment)", "https://www.federalreserve.gov/supervisionreg/srletters/SR2602a1.pdf")
p = doc.add_paragraph()
p.add_run("Legacy reference (superseded): ").bold = True
add_hyperlink(p, "SR 11-7, Guidance on Model Risk Management (April 4, 2011)", "https://www.federalreserve.gov/supervisionreg/srletters/sr1107.htm")
doc.add_paragraph(
    "The source list is limited to model-risk guidance addressed by this document. Applicable accounting, capital, stress-testing, CRE concentration, data, privacy, cybersecurity, fair-lending, and records-retention requirements require separate applicability analysis."
)

# Apply consistent section geometry/header/footer to every generated section.
for sec in doc.sections:
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.header.is_linked_to_previous = True
    sec.footer.is_linked_to_previous = True

doc.core_properties.title = "Commercial Real Estate Expected Loss Model Governance and Design"
doc.core_properties.subject = "SR 26-2 aligned governance and design for PD, LGD, EAD, scenarios, validation, and monitoring"
doc.core_properties.author = "CRE Expected Loss Project"
doc.core_properties.keywords = "CRE, expected loss, PD, LGD, EAD, model design"
doc.save(OUT)
print(OUT)
